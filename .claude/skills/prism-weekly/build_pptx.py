#!/usr/bin/env python3
"""prism-weekly · 슬라이드 PNG 들을 16:9 PPTX 로 묶는다.

사용: python3 build_pptx.py <출력.pptx> <png 디렉토리 | png 파일들...>
디렉토리를 주면 그 안의 slide-*.png 를 이름순으로 사용한다.
python-pptx 가상환경(.venv)이 없으면 자동으로 만든다.
"""
import os
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent
VENV_PY = ROOT / ".venv" / "bin" / "python3"

# PowerPoint 표준 16:9 (EMU)
SLIDE_W = 12192000
SLIDE_H = 6858000


def ensure_venv() -> None:
    print("bootstrapping python-pptx venv...", file=sys.stderr)
    subprocess.run([sys.executable, "-m", "venv", str(ROOT / ".venv")], check=True)
    subprocess.run([str(ROOT / ".venv" / "bin" / "pip"), "install", "-q", "python-pptx"], check=True)


def main() -> None:
    if len(sys.argv) < 3:
        sys.exit("usage: build_pptx.py <out.pptx> <png dir | png files...>")
    try:
        from pptx import Presentation
        from pptx.util import Emu
    except ImportError:
        if Path(sys.executable).resolve() == VENV_PY.resolve():
            sys.exit(".venv 에 python-pptx 가 없습니다. .venv 를 삭제하고 다시 실행하세요.")
        if not VENV_PY.exists():
            ensure_venv()
        os.execv(str(VENV_PY), [str(VENV_PY), __file__] + sys.argv[1:])

    out = Path(sys.argv[1]).resolve()
    args = [Path(a) for a in sys.argv[2:]]
    if len(args) == 1 and args[0].is_dir():
        pngs = sorted(args[0].glob("slide-*.png"))
    else:
        pngs = args
    if not pngs:
        sys.exit("PNG 가 없습니다.")

    prs = Presentation()
    prs.slide_width = Emu(SLIDE_W)
    prs.slide_height = Emu(SLIDE_H)
    blank = prs.slide_layouts[6]
    for p in pngs:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(p), 0, 0, width=prs.slide_width, height=prs.slide_height)
    prs.save(str(out))
    print(f"OK -> {out} ({out.stat().st_size // 1024} KB · {len(pngs)}장)", file=sys.stderr)


if __name__ == "__main__":
    main()

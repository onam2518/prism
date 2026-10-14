"""결과 → 키워드 실험실 양식 xlsx(100건씩 분할) + 전체 정리본 · 사용:
python3 deliver.py <input.json> <결과.jsonl>[,<재시도.jsonl>...] <양식.xlsx> <출력 접두 경로>
뒤에 오는 jsonl 이 같은 id 를 덮는다(재시도 결과 우선)."""
import json, re, shutil, sys, openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
MAX = 100                                            # keywordlab.MAX_ITEMS (업로드 1회 상한)
items = json.load(open(sys.argv[1]))
recs = {}
for f in sys.argv[2].split(','):
    for line in open(f):
        r = json.loads(line); recs[r['id']] = r
TPL, PREFIX = sys.argv[3], sys.argv[4]
name = lambda e: (e.get('name') if isinstance(e, dict) else str(e)).strip()
typed = lambda e: name(e) + (f"({e['type']})" if isinstance(e, dict) and e.get('type') else '')
cat = lambda c: ' / '.join(x for x in (c.get('tier1'), c.get('tier2')) if x) if isinstance(c, dict) else str(c)
missing = [it['id'] for it in items if it['id'] not in recs]
assert not missing, f'결과 없는 id: {missing[:5]}'
# ① 키워드 실험실 양식(콘텐츠 이름·리드문·엔티티·인텐트·카테고리 · 쉼표 구분 · 엔티티 이름만)
rows = [[it['title'][:300], recs[it['id']].get('summary', ''),
         ', '.join(name(e) for e in recs[it['id']].get('entities') or []),
         ', '.join(recs[it['id']].get('intent') or []),
         ', '.join(cat(c) for c in recs[it['id']].get('content_category') or [])] for it in items]
split = lambda v: [x.strip() for x in re.split(r'[,;|\n]', str(v or '')) if x.strip()]   # keywordlab.parse_excel 규칙
files = []
for i in range(0, len(rows), MAX):
    out = f'{PREFIX}_keywordlab_{i // MAX + 1}.xlsx'
    shutil.copy(TPL, out)
    wb = openpyxl.load_workbook(out); ws = wb.worksheets[0]
    ws.delete_rows(2, ws.max_row)
    for r in rows[i:i + MAX]:
        ws.append(r)
    wb.save(out)
    ok = sum(1 for r in rows[i:i + MAX] if r[1] and split(r[2]))
    files.append((out, len(rows[i:i + MAX]), ok))
# ② 전체 정리본(ID·타이틀·본문·리드문·인텐트·엔티티(유형)·카테고리)
wb = openpyxl.Workbook(); ws = wb.active; ws.title = '메타 추출'
ws.append(['ID', '타이틀', '본문', '리드문', '인텐트', '엔티티', '카테고리'])
for it in items:
    r = recs[it['id']]; st = r.get('meta_status') or {}
    ents = ' · '.join(typed(e) for e in r.get('entities') or [])
    if st.get('entities') not in ('success', 'no_value'):
        ents = '(추출 실패 · 재시도 소진)'
    ws.append([it['id'], it['title'], it['body'], r.get('summary', ''), ' · '.join(r.get('intent') or []),
               ents, ' · '.join(cat(c) for c in r.get('content_category') or [])])
for c in ws[1]:
    c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='E8EEF7')
for col, w in zip('ABCDEFG', (26, 40, 60, 50, 30, 50, 40)):
    ws.column_dimensions[col].width = w
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'B2'
wb.save(f'{PREFIX}_meta.xlsx')
for f in files:
    print('양식', f[0], f'{f[1]}건 · 업로드 반영 {f[2]}건')
print('전체', f'{PREFIX}_meta.xlsx')
print('엔티티 없음(업로드 제외)', [it['id'] for it in items if not recs[it['id']].get('entities')])
print('실패 상태', {it['id']: {k: v for k, v in (recs[it['id']].get('meta_status') or {}).items() if v not in ('success', 'no_value')}
                for it in items if any(v not in ('success', 'no_value') for v in (recs[it['id']].get('meta_status') or {}).values())})
